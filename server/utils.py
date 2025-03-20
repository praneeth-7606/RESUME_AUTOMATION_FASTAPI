import os
import json
import fitz  # PyMuPDF
import pandas as pd
from docx import Document
import io
from mistralai import Mistral
from fpdf import FPDF

def extract_text_from_file(file_path: str) -> str:
    """Extract text from various file formats."""
    file_extension = file_path.split('.')[-1].lower()
    text = ''

    if file_extension == 'pdf':
        doc = fitz.open(file_path)
        for page in doc:
            text += page.get_text("text")
    elif file_extension == 'docx':
        doc = Document(file_path)
        for paragraph in doc.paragraphs:
            text += paragraph.text + '\n'
    elif file_extension == 'txt':
        with open(file_path, 'r', encoding='utf-8') as file:
            text = file.read()
    else:
        raise ValueError("Unsupported file format. Use PDF, DOCX, or TXT.")

    return text

def load_skill_matrix(file_path: str) -> list:
    """Load and process skill matrix Excel file"""
    xls = pd.ExcelFile(file_path)
    sheets_data = []

    for sheet_name in xls.sheet_names[1:]:
        df = pd.read_excel(xls, sheet_name=sheet_name)
        records = df.to_dict(orient="records")
        sheets_data.append({"Sheet Name": sheet_name, "Data": records})

    return sheets_data

def LLM_CALL_1(resume_text: str) -> str:
    """Extract structured JSON from resume text"""
    client = Mistral(api_key=API_KEY)
    response = client.chat.complete(
        model="mistral-large-latest",
        messages=[{"role": "user", "content": f"Extract structured JSON from this resume:\n{resume_text}"}],
        response_format={"type": "json_object"}
    )
    return response.choices[0].message.content

def generate_resume_pdf(data: dict) -> str:
    """Generate a resume PDF"""
    pdf = FPDF()
    pdf.set_margins(15, 15, 15)
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 8, data["name"], ln=True, align="L")
    pdf.set_font("Arial", "", 11)
    pdf.multi_cell(0, 7, data["objective"])

    filename = os.path.join(OUTPUT_DIR, f"{data['name'].replace(' ', '_')}_Resume.pdf")
    pdf.output(filename)
    return filename
